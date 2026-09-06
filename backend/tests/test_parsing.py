import pytest

from app.services.ingestion.parsing import parse_file


def test_markdown_headings(tmp_path):
    p = tmp_path / "a.md"
    p.write_text("# 总览\nintro\n\n## 安装\nrun pip install\n", encoding="utf-8")
    blocks = parse_file(p, "text/markdown")
    assert blocks[0]["heading_path"] == "总览"
    assert blocks[1]["heading_path"] == "总览 > 安装"
    assert "pip install" in blocks[1]["text"]


def test_unsupported_type(tmp_path):
    p = tmp_path / "a.exe"
    p.write_bytes(b"x")
    with pytest.raises(ValueError):
        parse_file(p, "application/octet-stream")


def test_text_paragraphs(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("第一段。\n\n\n第二段。\n", encoding="utf-8")
    blocks = parse_file(p, "text/plain")
    assert [b["text"] for b in blocks] == ["第一段。", "第二段。"]
    assert all(b["page_no"] is None and b["heading_path"] == "" for b in blocks)


def test_pdf_roundtrip(tmp_path):
    import pymupdf

    p = tmp_path / "a.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "PDF第一页内容。这是首页的完整正文段落，用来提供足够的文本层。",
                     fontname="china-s")
    doc.new_page().insert_text((72, 72), "PDF第二页内容。次页正文同样保持足够的可提取字符数。",
                               fontname="china-s")
    doc.save(p)
    doc.close()

    blocks = parse_file(p, "application/pdf")
    assert any("PDF第一页内容" in b["text"] and b["page_no"] == 1 for b in blocks)
    assert any("PDF第二页内容" in b["text"] and b["page_no"] == 2 for b in blocks)


def test_pdf_table_to_markdown(tmp_path):
    """程序化生成含正文段落 + 网格表格的 PDF：表格转 Markdown 管道行，正文完整且不重复。"""
    import pymupdf

    p = tmp_path / "table.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "这是一段正文介绍，说明系统的整体架构与目标。",
                     fontname="china-s", fontsize=11)
    page.insert_text((72, 130), "下表列出各模块的职责与负责人。", fontname="china-s", fontsize=11)
    # 3 列 3 行网格表（画线触发 find_tables 的 lines_strict 检测）
    cols = [72, 172, 292, 412]
    rows = [160, 195, 230, 265]
    for x in cols:
        page.draw_line((x, rows[0]), (x, rows[-1]))
    for y in rows:
        page.draw_line((cols[0], y), (cols[-1], y))
    data = [["模块", "职责", "负责人"], ["解析", "PDF 转 MD", "张三"], ["检索", "混合召回", "李四"]]
    for r, row in enumerate(data):
        for c, cell in enumerate(row):
            page.insert_text((cols[c] + 6, rows[r] + 22), cell, fontname="china-s", fontsize=10)
    page.insert_text((72, 320), "表格之后还有一段总结性文字，用于验证正文完整性。",
                     fontname="china-s", fontsize=11)
    doc.save(p)

    blocks = parse_file(p, "application/pdf")
    # 表格块：Markdown 管道表行 + 分隔行，单元格齐全
    table_blocks = [b["text"] for b in blocks if b["text"].startswith("| ")]
    assert len(table_blocks) == 1
    lines = table_blocks[0].splitlines()
    assert lines[0].startswith("| 模块 | 职责 | 负责人 |")
    assert lines[1] == "| --- | --- | --- |"
    assert any("解析" in ln and "PDF 转 MD" in ln and "张三" in ln for ln in lines)
    assert all(b["page_no"] == 1 for b in blocks if b["text"].startswith("| "))
    # 正文完整：表格前后的段落都在
    joined = "\n".join(b["text"] for b in blocks)
    assert "这是一段正文介绍" in joined
    assert "表格之后还有一段总结性文字" in joined
    # 不重复：表格单元格文字不再作为普通文本块出现
    non_table = [b["text"] for b in blocks if not b["text"].startswith("| ")]
    assert all("张三" not in t and "李四" not in t for t in non_table)


def test_pdf_scanned_rejected(tmp_path):
    """纯图片 PDF（无可提取文本层）应明确报错而非产出空文档。"""
    import pymupdf

    p = tmp_path / "scan.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40))
    page.insert_image(pymupdf.Rect(72, 72, 112, 112), pixmap=pix)
    doc.save(p)

    with pytest.raises(ValueError, match="疑似扫描件 PDF"):
        parse_file(p, "application/pdf")


def test_docx_roundtrip(tmp_path):
    from docx import Document

    p = tmp_path / "a.docx"
    doc = Document()
    doc.add_heading("指南", level=1)
    doc.add_paragraph("安装说明正文。")
    doc.add_heading("配置", level=2)
    doc.add_paragraph("配置说明正文。")
    doc.save(p)

    blocks = parse_file(p, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert blocks[0]["text"] == "安装说明正文。"
    assert blocks[0]["heading_path"] == "指南"
    assert blocks[1]["text"] == "配置说明正文。"
    assert blocks[1]["heading_path"] == "指南 > 配置"
